#!/usr/bin/env python3
"""slides.html から PowerPoint（.pptx）を作る。

    python3 tools/build-pptx.py

各ページは全面1枚の画像。CSS で作った紙の縁やリーダー点をそのまま残すため、
図形に組み直さず見た目を優先している（その代わり PowerPoint 上で文字は編集できない）。
文字を直すときは slides.html を直して、これをもう一度流す。

発表メモは data-note から拾って PowerPoint のノートに入れるので、発表者ツールで読める。

必要なもの:
  - Google Chrome（ページの書き出しに使う）
  - python-pptx       pip3 install python-pptx
"""

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "slides.html"
OUTPUT = ROOT / "jev-lt-slides.pptx"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

# スライドの実寸。slides.html の #stage と合わせる。
STAGE_W, STAGE_H = 1280, 720
# 書き出しの倍率。投影しても粗く見えないよう 2 倍で焼く。
SCALE = 2

DECK_TITLE = "Jevで作る実用アプリ：家計簿のデータ設計から動くものになるまで"


def die(message: str) -> "typing.NoReturn":  # noqa: F821
    print(f"build-pptx: {message}", file=sys.stderr)
    sys.exit(1)


def strip_tags(html: str) -> str:
    text = re.sub(r"<[^>]+>", " ", html)
    for entity, char in [("&quot;", '"'), ("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">")]:
        text = text.replace(entity, char)
    return re.sub(r"\s+", " ", text).strip()


def main() -> None:
    if not SOURCE.exists():
        die(f"{SOURCE} がありません")
    if not Path(CHROME).exists():
        die("Google Chrome が見つかりません（ページの書き出しに使います）")
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        die("python-pptx がありません。pip3 install python-pptx")

    body = SOURCE.read_text(encoding="utf-8")

    sections = re.findall(r'<section class="slide"[\s\S]*?</section>', body)
    if not sections:
        die("slides.html にスライドが見つかりません")

    slides = []
    for index, section in enumerate(sections, start=1):
        head = re.search(r"<h[12][^>]*>([\s\S]*?)</h[12]>", section)
        eyebrow = re.search(r'<div class="eyebrow[^"]*">([\s\S]*?)</div>', section)
        note = re.search(r'data-note="([^"]*)"', section)
        title = strip_tags(head.group(1)) if head else ""
        if not title and eyebrow:
            title = strip_tags(eyebrow.group(1))
        slides.append({
            "n": index,
            "title": title or f"スライド {index}",
            "note": strip_tags(note.group(1)) if note else "",
        })

    print(f"スライド {len(slides)} ページを書き出します…")

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)

        # 書き出し用に手を入れる。
        #  - 入場のアニメーションを止める（途中の状態で焼けるのを防ぐ）
        #  - 操作の案内を消す
        #  - 残り時間は静止画だと 08:00 のまま固まって紛らわしいので隠す
        #  - 発表メモは PowerPoint のノートに入れるので、画面には出さない
        title_tag = re.search(r"<title>[\s\S]*?</title>", body)
        page = (
            '<!doctype html><html data-theme="dark"><head><meta charset="utf-8">'
            + (title_tag.group(0) if title_tag else "")
            + "</head><body>"
            + (body.replace(title_tag.group(0), "") if title_tag else body)
            + "</body></html>"
        ).replace(
            "</style>",
            "\n.slide.on{animation:none !important}\n"
            "#help{display:none}\n"
            "#clock{visibility:hidden}\n"
            "#notes{display:none !important}\n</style>",
        )
        export = tmp / "export.html"
        export.write_text(page, encoding="utf-8")

        shots = tmp / "png"
        shots.mkdir()
        for slide in slides:
            target = shots / f"s{slide['n']:02d}.png"
            subprocess.run(
                [
                    CHROME, "--headless", "--disable-gpu", "--no-sandbox",
                    f"--force-device-scale-factor={SCALE}",
                    "--virtual-time-budget=1800",
                    f"--window-size={STAGE_W},{STAGE_H}",
                    "--hide-scrollbars",
                    f"--screenshot={target}",
                    f"file://{export}#{slide['n']}",
                ],
                capture_output=True,
                check=False,
            )
            if not target.exists():
                die(f"{slide['n']} ページ目の書き出しに失敗しました")
            print(f"  {slide['n']:2d}/{len(slides)}  {slide['title'][:34]}")

        prs = Presentation()
        # 16:9。PowerPoint の既定のワイド画面と同じ寸法
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank = prs.slide_layouts[6]

        for slide in slides:
            page_ = prs.slides.add_slide(blank)
            page_.shapes.add_picture(
                str(shots / f"s{slide['n']:02d}.png"),
                0, 0, width=prs.slide_width, height=prs.slide_height,
            )
            # 画像だけだと読み上げソフトから中身が見えないので代替テキストを入れる
            page_.shapes[0]._element.nvPicPr.cNvPr.set("descr", slide["title"])

            notes = page_.notes_slide.notes_text_frame
            notes.text = slide["note"] or "（メモなし）"
            for paragraph in notes.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(14)

        prs.core_properties.title = DECK_TITLE
        prs.core_properties.subject = "LT（8分）"
        prs.core_properties.comments = "slides.html から書き出した画像。原本はそちら。"
        prs.save(str(OUTPUT))

    size_mb = OUTPUT.stat().st_size / 1024 / 1024
    print(f"\n保存しました: {OUTPUT.name}  ({len(slides)} ページ / {size_mb:.1f}MB)")


if __name__ == "__main__":
    main()
